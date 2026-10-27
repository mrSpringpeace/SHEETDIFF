"""Excel report: list Změny (všechny změny k filtrování) a list Souhrn."""
from __future__ import annotations

import os

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import __version__
from .compare import Result, cat_name
from .i18n import _

_HEAD_FILL = PatternFill("solid", fgColor="DDE4F0")
_HEAD_FONT = Font(bold=True)
_WRAP = Alignment(wrap_text=True, vertical="top")
_MAX = 32000  # Excel snese 32 767 znaků v buňce


def _text(v) -> str:
    s = ILLEGAL_CHARACTERS_RE.sub("", str(v))
    return s if len(s) <= _MAX else s[:_MAX] + " …"


def _put(ws, row, col, value, **style):
    c = ws.cell(row=row, column=col)
    c.value = _text(value)
    c.data_type = "s"   # text začínající = nesmí být vzorec
    for k, v in style.items():
        setattr(c, k, v)
    return c


def write(res: Result, path: str):
    wb = Workbook()
    ws = wb.active
    ws.title = _("Změny")
    head = [_(h) for h in ("Vrstva", "List", "Místo", "Změna", "Druh", "Staré (A)", "Nové (B)",
                           "Poznámka")]
    widths = [11, 22, 18, 30, 22, 60, 60, 40]
    for i, (h, w) in enumerate(zip(head, widths), 1):
        _put(ws, 1, i, h, font=_HEAD_FONT, fill=_HEAD_FILL)
        ws.column_dimensions[get_column_letter(i)].width = w
    row = 1
    for c in res.changes:
        row += 1
        cat = cat_name(c.cat)
        for i, v in enumerate([c.layer, c.sheet, c.loc, c.kind, cat, c.old, c.new, c.note], 1):
            _put(ws, row, i, v, alignment=_WRAP)
    ws.freeze_panes = "A2"
    if row > 1:
        ws.auto_filter.ref = f"A1:H{row}"

    ss = wb.create_sheet(_("Souhrn"))
    ss.column_dimensions["A"].width = 26
    ss.column_dimensions["B"].width = 100
    lines = [
        (_("Starý soubor (A)"), res.a.path),
        (_("Nový soubor (B)"), res.b.path),
        (_("Verdikt"), res.headline),
        (_("Návratový kód"), str(res.exit_code)),
        ("SHEETDIFF", __version__),
    ]
    lines += [("", f) for f in res.facts]
    lines.append(("", ""))
    lines += [(_("Šum"), f"{k} ({v})" if v not in ("", "1") else k) for k, v in res.noise]
    lines.append(("", ""))
    lines += [(_("Balíček"), f"{p.part} – {_(p.status)}" + (f": {p.detail}" if p.detail else ""))
              for p in res.parts]
    for i, (k, v) in enumerate(lines, 1):
        _put(ss, i, 1, k, font=_HEAD_FONT)
        _put(ss, i, 2, v, alignment=_WRAP)
    os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
    wb.save(path)
