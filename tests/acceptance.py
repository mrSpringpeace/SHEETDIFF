"""Akceptační testy SHEETDIFF nad vyrobenými sešity (vymyšlená data).

Spuštění z CODE_BASE:  python tests\\acceptance.py
Ukázkový rozpočet vyrábí tools/demo.py, další sešity a upravené varianty vznikají
v dočasné složce. Žádná vstupní data se nečtou.
"""
from __future__ import annotations

import os
import re
import sys
import tempfile
import time
import zipfile

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
sys.path.insert(0, os.path.join(ROOT, "tools"))
sys.stdout.reconfigure(encoding="utf-8", errors="replace")

import demo  # noqa: E402
from sheetdiff.compare import compare  # noqa: E402
from sheetdiff.report import render  # noqa: E402

TMP = tempfile.mkdtemp(prefix="sheetdiff_test_")
DEMO = demo.build(os.path.join(TMP, "demo_v1.xlsx"))


def variant(src: str, name: str, edits: dict[str, callable], level: int | None = None) -> str:
    """Kopie sešitu s upravenými částmi; edits: část -> funkce(bytes) -> bytes."""
    dst = os.path.join(TMP, name)
    with zipfile.ZipFile(src) as zi, zipfile.ZipFile(dst, "w") as zo:
        for info in zi.infolist():
            data = zi.read(info)
            if info.filename in edits:
                data = edits[info.filename](data)
            zo.writestr(info.filename, data, compress_type=zipfile.ZIP_DEFLATED,
                        compresslevel=level if level is not None else 6)
    return dst


def sub(pattern: bytes, repl: bytes, count=1):
    def f(data: bytes) -> bytes:
        out, n = re.subn(pattern, repl, data, count=count)
        assert n, f"vzor nenalezen: {pattern!r}"
        return out
    return f


failures = 0


def check(name, a, b, code, cats=None, kinds=()):
    global failures
    t = time.perf_counter()
    res = compare(a, b)
    html = render(res)
    dt = time.perf_counter() - t
    got = set(res.counts)
    ok = res.exit_code == code and (cats is None or got == set(cats)) and \
        all(any(k in c.kind for c in res.changes) for k in kinds)
    failures += not ok
    print(f"[{'OK ' if ok else 'CHYBA'}] {name}: kód {res.exit_code} (čekáno {code}), "
          f"{dt:.1f} s, {len(html) // 1024} kB report")
    print(f"        {res.headline}  {res.facts}")
    if not ok:
        for c in res.changes[:10]:
            print("        ", c.layer, c.sheet, c.loc, c.kind, c.old[:50], "->", c.new[:50], c.note)
        for p in res.parts:
            print("         část:", p.part, p.status, p.detail)


DRAWING, CHART, SHEET = "xl/drawings/drawing1.xml", "xl/charts/chart1.xml", "xl/worksheets/sheet1.xml"


def demo_with(*changes):
    return demo.build(os.path.join(TMP, "demo_" + "_".join(changes) + ".xlsx"), changes)


check("T1 stejný soubor dvakrát", DEMO, DEMO, 0)
check("T2 jiná komprese", DEMO, variant(DEMO, "t2.xlsx", {}, level=1), 1)
check("T3 text v textovém poli", DEMO,
      variant(DEMO, "t3.xlsx", {DRAWING: sub(rb"Prices exclude", b"Prices include")}), 2, {"shape_text"})
check("T4 jen texty buněk", DEMO, demo_with("text"), 2, {"text"})
check("T5 změněná řada grafu", DEMO,
      variant(DEMO, "t5.xlsx", {CHART: sub(rb"\$D\$2:\$D\$8", rb"$D$2:$D$7")}), 2, {"chart"}, ["Graf: Řada"])
if zipfile.ZipFile(DEMO).namelist().count("xl/media/image1.png"):
    check("T6 vyměněný obrázek", DEMO,
          variant(DEMO, "t6.xlsx", {"xl/media/image1.png": lambda d: d[:-100] + bytes(100)}),
          2, {"image"}, ["Obrázek vyměněn"])
else:
    print("[--] T6 vynechán: bez Pillow nemá ukázkový sešit obrázek")
check("vzorec změněn", DEMO, demo_with("formula"), 2, {"formula"}, ["Vzorec změněn"])
check("číslo změněno", DEMO, demo_with("number"), 2, {"number"})
check("jen aktivní buňka (šum)", DEMO,
      variant(DEMO, "tv.xlsx", {SHEET: sub(rb'activeCell="[A-Z]+\d+"', b'activeCell="Z99"')}), 1)
check("neznámá část (záchranná síť)", DEMO,
      variant(DEMO, "tu.xlsx", {"xl/theme/theme1.xml": sub(rb'name="[^"]*"', b'name="Other"')}),
      2, {"other"})

# všechny změny ukázkového rozpočtu najednou: každá právě jednou, vložený řádek na správném místě
full = compare(DEMO, demo_with(*demo.ALL))
expected = {"struct": 2, "text": 1, "number": 1, "formula": 1, "format": 1, "rules": 2, "comment": 1,
            "chart": 1, "shape_text": 1}
if "xl/media/image1.png" in zipfile.ZipFile(DEMO).namelist():
    expected["image"] = 1
ins = [c for c in full.changes if c.kind == "Vložen řádek"]
ok = dict(full.counts) == expected and len(ins) == 1 and ins[0].loc == "4"
failures += not ok
print(f"[{'OK ' if ok else 'CHYBA'}] ukázkový rozpočet v1 → v2: {sum(full.counts.values())} změn")
if not ok:
    print("         ", dict(full.counts))
    for c in full.changes:
        print("         ", c.layer, c.loc, c.kind, c.note)

# ---------------------------------------------------------------- vyrobené sešity (openpyxl)

from openpyxl import Workbook, load_workbook  # noqa: E402
from openpyxl.comments import Comment  # noqa: E402
from openpyxl.formatting.rule import CellIsRule  # noqa: E402
from openpyxl.worksheet.datavalidation import DataValidation  # noqa: E402
from openpyxl.styles import Font, PatternFill  # noqa: E402
from openpyxl.utils import get_column_letter as L  # noqa: E402


def book(name, insert_row=False, delete_row=False, insert_col=False, calc="Calc",
         bold=False, merge=False, hide=False, comments=(), cf=None, dv=None, link=None):
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    rows = [(f"Položka {i}", i * 10) for i in range(1, 31)]
    if insert_row:
        rows.insert(10, ("Nová položka", 999))
    if delete_row:
        rows.pop(4)
    v, f = (3, 4) if insert_col else (2, 3)
    for r, (lab, val) in enumerate(rows, 1):
        ws.cell(r, 1, lab)
        ws.cell(r, v, val)
        ws.cell(r, f, f"={L(v)}{r}*2")
    n = len(rows)
    ws.cell(n + 2, f, f"=SUM({L(f)}1:{L(f)}{n})")
    ws.cell(n + 3, 1, f"={calc}!A1")
    w2 = wb.create_sheet(calc)
    w2["A1"] = 42
    w2["B1"] = f"=Data!{L(v)}2"
    if bold:
        for r in (3, 4, 5):
            ws.cell(r, v).font = Font(bold=True)
            ws.cell(r, v).fill = PatternFill("solid", fgColor="FFFF00")
    if merge:
        ws.merge_cells("H1:I1")
    if hide:
        ws.row_dimensions[3].hidden = True
    for addr, text in comments:
        ws[addr].comment = Comment(text, "JS")
    if cf:
        ws.conditional_formatting.add(f"{L(v)}1:{L(v)}{n}", CellIsRule(
            operator="greaterThan", formula=[cf], fill=PatternFill("solid", bgColor="FFC7CE")))
    if dv:
        d = DataValidation(type="list", formula1=f'"{dv}"', allow_blank=True)
        d.add("E1:E5")
        ws.add_data_validation(d)
    if link:
        ws["A1"].hyperlink = link
    path = os.path.join(TMP, name)
    wb.save(path)
    return path


def _with_vba(src, name, payload):
    dst = os.path.join(TMP, name)
    with zipfile.ZipFile(src) as zi, zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zo:
        for info in zi.infolist():
            zo.writestr(info.filename, zi.read(info))
        zo.writestr("xl/vbaProject.bin", payload)
    return dst


def check_n(name, a, b, code, cats, kinds, n):
    res = compare(a, b)
    real = [c for c in res.changes if c.cat in cats]
    global failures
    ok = res.exit_code == code and set(res.counts) == set(cats) and len(real) == n and \
        all(any(k in c.kind for c in res.changes) for k in kinds)
    failures += not ok
    print(f"[{'OK ' if ok else 'CHYBA'}] {name}: kód {res.exit_code}, {len(real)} změn (čekáno {n})")
    print(f"        {res.headline}")
    for c in res.changes if not ok else real[:3]:
        print("        ", c.layer, c.sheet, c.loc, "|", c.kind, "|", c.old[:60].replace("\n", " / "),
              "->", c.new[:60].replace("\n", " / "), "|", c.note)
    return res


base = book("base.xlsx")
check_n("T7 vložený řádek", base, book("t7.xlsx", insert_row=True), 2, {"struct"}, ["Vložen řádek"], 1)
check_n("T7b vložený prázdný sloupec", base, book("t7b.xlsx", insert_col=True), 2, {"struct"},
        ["Vložen sloupec"], 1)
check_n("T7c smazaný řádek", base, book("t7c.xlsx", delete_row=True), 2, {"struct"}, ["Smazán řádek"], 1)
check_n("T8 přejmenovaný list", base, book("t8.xlsx", calc="Výpočet"), 2, {"struct"},
        ["přejmenován"], 1)
check_n("formát buněk", base, book("tfmt.xlsx", bold=True), 2, {"format"}, ["Formát změněn"], 1)
check_n("sloučené buňky", base, book("tmrg.xlsx", merge=True), 2, {"format"}, ["sloučeny"], 1)
check_n("skrytý řádek", base, book("thid.xlsx", hide=True), 2, {"format"}, ["skryty"], 1)
cm_a = book("tcm_a.xlsx", comments=[("A2", "Pozor"), ("A4", "Starý text")])
res = check_n("komentáře", cm_a, book("tcm_b.xlsx", comments=[("A4", "Nový text"), ("A6", "Přidán")]),
              2, {"comment"}, ["Komentář změněn", "přidán", "odebrán"], 3)

cf_a = book("tcf_a.xlsx", cf="100", dv="a,b", link="https://example.com/a")
check_n("podmíněný formát přidán", base, book("tcf_add.xlsx", cf="100"), 2, {"rules"}, ["pravidlo přidáno"], 1)
check_n("podmíněný formát změněn", cf_a, book("tcf_chg.xlsx", cf="200", dv="a,b", link="https://example.com/a"),
        2, {"rules"}, ["pravidlo změněno"], 1)
check_n("ověření dat změněno", cf_a, book("tdv.xlsx", cf="100", dv="a,b,c", link="https://example.com/a"),
        2, {"rules"}, ["Ověření dat"], 1)
check_n("odkaz změněn", cf_a, book("tln.xlsx", cf="100", dv="a,b", link="https://example.com/b"),
        2, {"link"}, ["Odkaz změněn"], 1)
check_n("vložený řádek + pravidla beze změny", cf_a,
        book("tcf_ins.xlsx", cf="100", dv="a,b", link="https://example.com/a", insert_row=True),
        2, {"struct"}, ["Vložen řádek"], 1)
check_n("makra VBA", variant(base, "vba_a.xlsx", {}, level=6) and _with_vba(base, "vba_a.xlsx", b"A" * 100),
        _with_vba(base, "vba_b.xlsx", b"B" * 100), 2, {"vba"}, ["Makra VBA změněna"], 1)

from sheetdiff.report_xlsx import write  # noqa: E402
xp = os.path.join(TMP, "report.xlsx")
write(res, xp)
ws = load_workbook(xp)["Změny"]
ok = ws.max_row == len(res.changes) + 1 and ws["D2"].value
failures += not ok
print(f"[{'OK ' if ok else 'CHYBA'}] Excel report: {ws.max_row - 1} řádků změn")

# ---------------------------------------------------------------- angličtina

from sheetdiff.i18n import set_lang  # noqa: E402

CZ = re.compile(r"[ěščřžýůťďňĚŠČŘŽÝŮŤĎŇáíéú]")
set_lang("en")
try:
    leaks = []
    en_cases = [
        (DEMO, demo_with(*demo.ALL)),
        (DEMO, demo_with("text")),
        (base, book("en_ins.xlsx", insert_row=True)),
        (base, book("en_fmt.xlsx", bold=True, merge=True, hide=True)),
        (cf_a, book("en_cf.xlsx", cf="200", dv="a,b,c", link="https://example.com/b")),
        (cm_a, book("en_cm.xlsx", comments=[("A4", "Nový text")])),
        (base, book("en_t8.xlsx", calc="Výpočet")),
        (DEMO, os.path.join(TMP, "t5.xlsx")),
        (DEMO, os.path.join(TMP, "t3.xlsx")),
        (DEMO, os.path.join(TMP, "tu.xlsx")),
        (DEMO, os.path.join(TMP, "t2.xlsx")),
    ]
    for a, b in en_cases:
        res = compare(a, b)
        labels = [res.headline, *res.facts, *(k for k, _v in res.noise)]
        labels += [x for c in res.changes for x in (c.layer, c.kind)]
        labels += [c.note for c in res.changes if c.cat in ("struct", "shape", "image", "vba")]
        labels += [p.detail for p in res.parts]
        page = render(res)
        page = re.sub(r"const ROWS = .*?;\n", "", page)  # data ze sešitů
        page = re.sub(r"<td class=mono>.*?</td>|<div class=\"(name|path)\">.*?</div>|data-sheet=\"[^\"]*\"|"
                      r"<th scope=row>.*?</th>|<title>.*?</title>", "", page)
        page = re.sub(r"Listy beze změny: .*?</p>|Unchanged sheets: .*?</p>", "", page)
        page = re.sub(r"<option value=\"[^\"]*\">[^<]*</option>", "", page)
        leaks += [x for x in labels if CZ.search(x) and not any(s in x for s in res.sheets)]
        leaks += [m.group(0) for m in re.finditer(r".{30}" + CZ.pattern + r".{30}", page)]
    ok = not leaks
    failures += not ok
    print(f"[{'OK ' if ok else 'CHYBA'}] angličtina: {len(en_cases)} porovnání bez českých popisků")
    for x in leaks[:15]:
        print("         ", repr(x))
    r_en = compare(base, book("en_ins2.xlsx", insert_row=True))
    ok = r_en.changes[0].kind == "Row inserted" and r_en.headline.startswith("Only the structure")
    failures += not ok
    print(f"[{'OK ' if ok else 'CHYBA'}] angličtina: {r_en.changes[0].kind} / {r_en.headline}")
finally:
    set_lang("cs")

# ---------------------------------------------------------------- okno a přetažení

if sys.platform == "win32":
    import subprocess  # noqa: E402
    try:
        p = subprocess.run([sys.executable, os.path.join(ROOT, "tests", "gui_drop.py")],
                           capture_output=True, text=True, encoding="utf-8", timeout=90)
        out = (p.stdout + p.stderr).strip()
    except subprocess.TimeoutExpired:
        out = "časový limit vypršel (okno zamrzlo)"
    ok = out.endswith("GUI OK")
    failures += not ok
    print(f"[{'OK ' if ok else 'CHYBA'}] okno: přetažení souborů frontou zpráv")
    if not ok:
        print("         ", out[-1500:])

# ---------------------------------------------------------------- verze

from sheetdiff import __version__  # noqa: E402

changelog = open(os.path.join(ROOT, "CHANGELOG.md"), encoding="utf-8").read()
ok = re.fullmatch(r"\d+\.\d+\.\d+", __version__) is not None and f"\n## {__version__} (" in changelog
failures += not ok
print(f"[{'OK ' if ok else 'CHYBA'}] verze {__version__} má záznam v CHANGELOG.md")

print(f"\nVarianty v {TMP}")
print("VŠE OK" if not failures else f"SELHALO: {failures}")
sys.exit(1 if failures else 0)
