"""Ukázkové sešity s vymyšlenými daty (rozpočet kanceláře) pro testy, README a screenshot.

    python tools/demo.py examples        -> examples/office_budget_v1.xlsx a _v2.xlsx

build(path, changes) vytvoří verzi 1; každá položka v `changes` přidá jednu změnu verze 2:
    text, number, formula, insert_row, format, comment, cf, dv, textbox, chart_title, image, rename
Obrázek potřebuje Pillow; bez něj se sešit vytvoří bez obrázku.
"""
from __future__ import annotations

import io
import os
import sys
import zipfile

from lxml import etree
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ALL = ("text", "number", "formula", "insert_row", "format", "comment", "cf", "dv",
       "textbox", "chart_title", "image", "rename")

ITEMS = [
    ("Laptops", 12, 1150),
    ("Monitors", 12, 210),
    ("Office chairs", 12, 340),
    ("Desks", 12, 420),
    ("Software licences", 12, 95),
    ("Training", 1, 2400),
    ("Travel", 1, 3600),
]


def _logo(color: str) -> io.BytesIO | None:
    try:
        from PIL import Image, ImageDraw
    except ImportError:
        return None
    img = Image.new("RGB", (120, 60), "white")
    d = ImageDraw.Draw(img)
    d.rounded_rectangle((4, 4, 116, 56), radius=10, fill=color)
    d.rectangle((20, 24, 100, 36), fill="white")
    buf = io.BytesIO()
    img.save(buf, "PNG")
    buf.seek(0)
    return buf


def build(path: str, changes=()) -> str:
    ch = set(changes)
    wb = Workbook()
    ws = wb.active
    ws.title = "Budget"
    items = [list(x) for x in ITEMS]
    if "text" in ch:
        items[2][0] = "Ergonomic office chairs"
    if "number" in ch:
        items[1][2] = 189
    if "insert_row" in ch:
        items.insert(2, ["Docking stations", 12, 160])

    ws.append(["Item", "Qty", "Unit price", "Total"])
    for r, (name, qty, price) in enumerate(items, start=2):
        ws.append([name, qty, price, f"=B{r}*C{r}"])
    last = len(items) + 1
    ws.cell(last + 2, 3, "Subtotal")
    ws.cell(last + 2, 4, f"=SUM(D2:D{last})")
    ws.cell(last + 3, 3, "Reserve")
    ws.cell(last + 3, 4, f"=D{last + 2}*{'0.1' if 'formula' in ch else '0.05'}")
    ws.cell(last + 4, 3, "Grand total")
    ws.cell(last + 4, 4, f"=D{last + 2}+D{last + 3}")

    head = PatternFill("solid", fgColor="2F5597" if "format" in ch else "D9E2F3")
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF" if "format" in ch else "000000")
        c.fill = head
    for row in ws.iter_rows(min_row=2, min_col=3, max_col=4):
        for c in row:
            c.number_format = "#,##0"
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["C"].width = 12
    ws.column_dimensions["D"].width = 12

    ws["C2"].comment = Comment("Quote from supplier B, valid until March" if "comment" in ch
                               else "Quote from supplier A", "Reviewer")
    limit = "8000" if "cf" in ch else "5000"
    ws.conditional_formatting.add(f"D2:D{last}", CellIsRule(
        operator="greaterThan", formula=[limit], fill=PatternFill("solid", bgColor="FCE4D6")))
    if "dv" in ch:
        dv = DataValidation(type="whole", operator="between", formula1="0", formula2="100")
        dv.add(f"B2:B{last}")
        ws.add_data_validation(dv)

    chart = BarChart()
    chart.title = "Cost by item" if "chart_title" in ch else "Budget by item"
    chart.add_data(Reference(ws, min_col=4, min_row=1, max_row=last), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=2, max_row=last))
    chart.width, chart.height = 14, 7
    ws.add_chart(chart, "F2")
    logo = _logo("#2f8a4f" if "image" in ch else "#2357c6")
    if logo is not None:
        from openpyxl.drawing.image import Image
        ws.add_image(Image(logo), "F18" if "insert_row" in ch else "F17")

    notes = wb.create_sheet("Read me" if "rename" in ch else "Notes")
    notes["A1"] = "Office budget, invented example data for SHEETDIFF."
    notes["A2"] = "All prices in EUR."

    wb.save(path)
    _add_textbox(path, "Draft budget for the new office. Prices include delivery."
                 if "textbox" in ch else "Draft budget for the new office. Prices exclude delivery.",
                 row=17 if "insert_row" in ch else 16)
    return path


_XDR = "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing"
_A = "http://schemas.openxmlformats.org/drawingml/2006/main"


def _add_textbox(path: str, text: str, row: int = 16):
    """openpyxl textová pole neumí, proto se doplní přímo do kresby prvního listu."""
    with zipfile.ZipFile(path) as z:
        parts = {i.filename: z.read(i) for i in z.infolist()}
    name = "xl/drawings/drawing1.xml"
    root = etree.fromstring(parts[name])
    anchor = etree.fromstring(f'''
<twoCellAnchor xmlns="{_XDR}" xmlns:a="{_A}">
  <from><col>10</col><colOff>0</colOff><row>{row}</row><rowOff>0</rowOff></from>
  <to><col>14</col><colOff>0</colOff><row>{row + 4}</row><rowOff>0</rowOff></to>
  <sp macro="" textlink="">
    <nvSpPr><cNvPr id="20" name="TextBox 1"/><cNvSpPr txBox="1"/></nvSpPr>
    <spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="2600000" cy="760000"/></a:xfrm>
      <a:prstGeom prst="rect"><a:avLst/></a:prstGeom>
      <a:solidFill><a:srgbClr val="FFF8E1"/></a:solidFill>
      <a:ln><a:solidFill><a:srgbClr val="C9A227"/></a:solidFill></a:ln></spPr>
    <txBody><a:bodyPr wrap="square"/><a:lstStyle/>
      <a:p><a:r><a:rPr lang="en-US" sz="1100"/><a:t>{text}</a:t></a:r></a:p></txBody>
  </sp>
  <clientData/>
</twoCellAnchor>''')
    root.append(anchor)
    parts[name] = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for n, data in parts.items():
            z.writestr(n, data)


if __name__ == "__main__":
    out = sys.argv[1] if len(sys.argv) > 1 else "examples"
    os.makedirs(out, exist_ok=True)
    print(build(os.path.join(out, "office_budget_v1.xlsx")))
    print(build(os.path.join(out, "office_budget_v2.xlsx"), ALL))
